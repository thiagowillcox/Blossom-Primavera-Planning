import json
import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from .hierarchy import activity_hierarchy
from .source_audit import reconcile_ld


def serial(value):
    if isinstance(value,Decimal):
        return str(value)
    raise TypeError(type(value).__name__)


def save(tables, records, metadata, directory):
    directory.mkdir(parents=True,exist_ok=True)
    hierarchy = activity_hierarchy(tables)
    source_checks=reconcile_ld(records,tables['Auditoria_HH'])
    (directory/'conferencia_ld.json').write_text(json.dumps(source_checks,ensure_ascii=False,indent=2,default=serial)+'\n',encoding='utf-8')
    payload={'conferencia_ld':source_checks,'linhas_cronograma': hierarchy, 'metadados': metadata,'tabelas':tables,'registros':records}
    (directory/'resultado.json').write_text(json.dumps(payload,ensure_ascii=False,sort_keys=True,indent=2,default=serial)+'\n', encoding='utf-8')
    (directory/'pendencias.json').write_text(json.dumps(tables['Pendencias'],ensure_ascii=False,indent=2,default=serial)+'\n', encoding='utf-8')
    (directory/'matriz_regras.json').write_text(json.dumps(tables['Matriz_Regras'],ensure_ascii=False,sort_keys=True,indent=2,default=serial)+'\n', encoding='utf-8')
    workbook=openpyxl.Workbook();workbook.remove(workbook.active)
    for name, rows in tables.items():
        if name == 'Atividades':
            rows = hierarchy
        sheet=workbook.create_sheet(name)
        keys=list(dict.fromkeys(k for row in rows for k in row)) or ['status']
        sheet.append(keys)
        for row in rows:
            values=[]
            for key in keys:
                v=row.get(key)
                if isinstance(v,Decimal):v=float(v)
                if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False,default=serial)
                values.append(v)
            sheet.append(values)
            # Source text and formulas are exported as inert strings.
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
        sheet.freeze_panes='A2'
        if name != 'Atividades':
            sheet.auto_filter.ref=sheet.dimensions
        else:
            # Excel filters/sorts can separate children from summary rows.
            sheet.sheet_properties.outlinePr.summaryBelow = False
            sheet.sheet_view.showOutlineSymbols = True
            name_col = keys.index('nome')+1
            for index, row in enumerate(rows, 2):
                level = row['nivel']
                sheet.row_dimensions[index].outlineLevel = min(level, 7)
                sheet.cell(index, name_col).alignment = Alignment(indent=min(level, 250))
                if row['tipo_linha'] != 'atividade':
                    for cell in sheet[index]:
                        cell.font = Font(bold=True, color='FFFFFF' if level == 0 else '183C30')
                        cell.fill = PatternFill('solid', fgColor='183C30' if level == 0 else 'E9E2E6')
        for cell in sheet[1]:
            cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='70304A')
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col[:100])+2))
    workbook.properties.created=workbook.properties.modified=datetime(2000,1,1)
    temporary=directory/'_temp.xlsx';workbook.save(temporary)
    # Normalize ZIP timestamps too: same inputs/config produce byte-identical XLSX.
    with ZipFile(temporary) as src, ZipFile(directory/'pre_cronograma.xlsx','w',ZIP_DEFLATED) as dst:
        for info in src.infolist():
            info.date_time=(2000,1,1,0,0,0)
            data=src.read(info.filename)
            if info.filename == 'docProps/core.xml':
                data=re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)', rb'\g<1>2000-01-01T00:00:00Z\g<2>', data)
            dst.writestr(info,data)
    temporary.unlink()
