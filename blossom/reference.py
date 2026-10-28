"""Reprocess an audited Wardi export as an explicit planning source.

This is not generation from a new LD: structure, IDs and approved network are
imported; effort comes from preserved LD caches; dates are recalculated.
"""
from io import BytesIO
import hashlib
import json
import re
from decimal import Decimal
import openpyxl
from .common import number
from .schedule import estimate


def reprocess_reference(content):
    workbook=openpyxl.load_workbook(BytesIO(content),data_only=False)
    required=['Cronograma','WBS','Atividades','Atribuicoes_Funcoes','Relacionamentos','Matriz_Regras','Parametros','Pendencias','Auditoria_HH']
    if any(name not in workbook.sheetnames for name in required):
        raise ValueError('Referência sem as nove abas auditadas exigidas pelo formato Wardi.')
    tables={};positions={}
    for name in required:
        sheet=workbook[name];headers=[c.value for c in sheet[1]]
        tables[name]=[];positions[name]={key:i+1 for i,key in enumerate(headers)}
        for row in sheet.iter_rows(min_row=2):
            if any(c.value is not None for c in row):
                tables[name].append((row[0].row,dict(zip(headers,[c.value for c in row]))))
    for name,field in [('WBS','id'),('Cronograma','ID'),('Atividades','id'),('Atribuicoes_Funcoes','id'),('Relacionamentos','id')]:
        ids=[r[field] for i,r in tables[name]]
        if len(ids)!=len(set(ids)) or any(v is None for v in ids):raise ValueError('IDs inválidos ou duplicados em '+name)
    params={r['parametro']:r['valor'] for index,r in tables['Parametros']}
    daily=number(params.get('hh_dia'))
    if daily is None or daily<=0:raise ValueError('Jornada de referência inválida.')
    fractions=[Decimal(str(v)) for v in json.loads(params['percentuais'])]
    if fractions!=list(map(Decimal,['.8','.1','.2','.1'])):raise ValueError('Percentuais incompatíveis com 80/10/20/10.')
    client=json.loads(params['cliente_dias'])
    for stage in ('2','4'):
        if number(client.get(stage)) is None:raise ValueError('Prazo do cliente ausente na referência.')
    original={};changes=[];acts={};source_rows={};record_types={}
    for index,r in tables['Auditoria_HH']:
        if r['funcao'] is None:continue
        match=re.fullmatch(r"='CPU '!([E-M])2",r['funcao'])
        if not match:raise ValueError('Função da referência sem correspondência explícita com H:P da LD.')
        cache=json.loads(r['cache'])
        if len(cache)<18:raise ValueError('Cache de auditoria incompleto.')
        if r.get('fracao_registro')!=1:raise ValueError('Rateio fracionário exige a regra original; reprocessamento não suportado.')
        role_index=ord(match.group(1))-ord('E')+7
        unit=cache[5];value=number(cache[role_index])
        if value is None:raise ValueError('HH de função ausentes no cache da referência.')
        multiplier=Decimal(1) if unit=='HH' else number(cache[6])
        if multiplier is None:raise ValueError('A1 ausente na auditoria da referência.')
        hours=value*multiplier
        key=(r['registro'],r['funcao'])
        if key in original:raise ValueError('Auditoria duplicada por registro/função.')
        if r['hh_original'] is None or abs(hours-number(r['hh_original']))>Decimal('.00000001'):
            raise ValueError('HH originais da auditoria divergem do cache LD preservado.')
        original[key]=hours
        if r['registro'] in source_rows and source_rows[r['registro']]!=cache:
            raise ValueError('Caches conflitantes para o mesmo registro.')
        source_rows[r['registro']]=cache

    def write(name,index,row,field,value):
        if field not in positions[name]:raise ValueError(f'Campo obrigatório ausente: {name}.{field}')
        before=row.get(field)
        numeric=isinstance(before,(int,float,Decimal)) and not isinstance(before,bool) and isinstance(value,(int,float,Decimal))
        equal=abs(Decimal(str(before))-Decimal(str(value)))<=Decimal('.00000001') if numeric else before==value
        if not equal:changes.append({'aba':name,'linha':index,'campo':field,'anterior':before,'recalculado':str(value)})
        row[field]=value
        workbook[name].cell(index,positions[name][field],float(value) if isinstance(value,Decimal) else value)

    for index,r in tables['Atividades']:
        aid=r['id'];rid=r['registro']
        if aid in acts:raise ValueError('Atividade duplicada na referência.')
        if rid not in source_rows:raise ValueError('Atividade sem cache auditado.')
        auxiliary=r['tipo']=='auxiliar';stage=r['etapa']
        if not auxiliary and stage not in (1,2,3,4):raise ValueError('Etapa documental inválida.')
        if rid in record_types and record_types[rid]!=auxiliary:raise ValueError('Registro com tipos conflitantes.')
        record_types[rid]=auxiliary
        fraction=Decimal(1) if auxiliary else fractions[stage-1]
        hours=sum((h for (record,role),h in original.items() if record==rid),Decimal(0))*fraction
        days=number(client[str(stage)]) if not auxiliary and stage in (2,4) else hours/daily
        write('Atividades',index,r,'hh',hours);write('Atividades',index,r,'duracao_dias',days)
        write('Atividades',index,r,'quantidade',source_rows[rid][4]);write('Atividades',index,r,'a1',source_rows[rid][6])
        acts[aid]=dict(id=aid,registro=rid,hh=hours,duracao_dias=days,fraction=fraction)
    groups={rid:[] for rid in record_types}
    for index,r in tables['Atividades']:groups[r['registro']].append(r)
    for rid,group in groups.items():
        if record_types[rid]:
            if len(group)!=1:raise ValueError('Auxiliar com múltiplas atividades.')
        elif sorted(r['etapa'] for r in group)!=[1,2,3,4]:raise ValueError('Ciclo documental incompleto.')
    assigned={};seen=set()
    for index,r in tables['Atribuicoes_Funcoes']:
        aid=r['atividade'];key=(aid,r['funcao'])
        if aid not in acts or key in seen:raise ValueError('Atribuição órfã ou duplicada.')
        seen.add(key);a=acts[aid]
        if r['registro']!=a['registro'] or (r['registro'],r['funcao']) not in original:raise ValueError('Atribuição sem original auditado.')
        hours=original[(r['registro'],r['funcao'])]*a['fraction']
        write('Atribuicoes_Funcoes',index,r,'hh',hours)
        assigned[(r['registro'],r['funcao'])]=assigned.get((r['registro'],r['funcao']),Decimal(0))+hours
    for aid,a in acts.items():
        expected={(aid,role) for (rid,role),h in original.items() if rid==a['registro'] and h>0}
        if not expected.issubset(seen):raise ValueError('Função auditada sem atribuição de etapa.')
    links=[]
    for index,r in tables['Relacionamentos']:
        if r['predecessora'] not in acts or r['sucessora'] not in acts:raise ValueError('Vínculo com atividade inexistente.')
        if r.get('ativo_no_calculo') is True:
            if r['tipo']!='FS':raise ValueError('Relação ativa diferente de FS não suportada.')
            links.append(dict(predecessora=r['predecessora'],sucessora=r['sucessora'],defasagem_dias=0))
    estimate(list(acts.values()),links)
    for index,r in tables['Atividades']:
        minimum=r.get('inicio_minimo_dia_util')
        if minimum not in (None,0):raise ValueError('Restrição de início exige política original; reprocessamento não suportado.')
        for field,key in [('inicio_dia_util','inicio_dia'),('fim_dia_util','fim_dia')]:write('Atividades',index,r,field,acts[r['id']][key])
    for index,r in tables['Auditoria_HH']:
        if r['funcao'] is None:continue
        key=(r['registro'],r['funcao']);base=original[key]
        expected=base*(Decimal(1) if record_types[r['registro']] else Decimal('1.2'))
        actual=assigned.get(key,Decimal(0))
        if actual!=expected:raise ValueError('Falha de conservação de horas recalculadas.')
        for field,value in [('hh_original',base),('hh_planejada',actual),('hh_esperada',expected),('diferenca',actual-expected)]:write('Auditoria_HH',index,r,field,value)
    for index,r in tables['Cronograma']:
        if r['ID'] in acts:
            a=acts[r['ID']]
            for field,key in [('HH','hh'),('Duracao_dias','duracao_dias'),('Inicio_dia_util','inicio_dia'),('Fim_dia_util','fim_dia')]:write('Cronograma',index,r,field,a[key])
    buffer=BytesIO();workbook.save(buffer);workbook.close()
    summary={'sha256_referencia':hashlib.sha256(content).hexdigest(),'tolerancia_numerica':'0.00000001','fonte':'cache LD e rede importados da referência; não é geração a partir de nova LD',
             'atividades':len(acts),'relacionamentos_ativos':len(links),'hh_planejadas':str(sum((a['hh'] for a in acts.values()),Decimal(0))),
             'alteracoes_relevantes':len(changes),'diferencas':changes,
             'observacao':'Duração de produção = HH/jornada (equivalente homem-dia); condições externas e metadados são preservados da referência.'}
    return buffer.getvalue(),summary
