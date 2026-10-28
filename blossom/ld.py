"""Leitor do layout LD Blossom; nunca executa macros ou fórmulas arbitrárias."""
import json
import re
from decimal import Decimal
import openpyxl
from .common import stable_id, number, issue, clean_text
from .wbs import node

CODES = {'DC', 'DB', 'DD', 'AC', 'AB', 'AD'}
FORMATS = {'A4': Decimal('0.125'), 'A3': Decimal('0.25'),
           'A2': Decimal('0.5'), 'A1': Decimal(1), 'A0': Decimal(2)}


def read_ld(path, identities=None):
    raw = openpyxl.load_workbook(path, data_only=False, read_only=True)
    cached = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = [name for name in raw.sheetnames if clean_text(name).casefold() == 'ld']
    if len(sheets) != 1:
        raw.close(); cached.close()
        raise ValueError('Não foi possível identificar uma única aba LD. Confira o nome da aba no Excel.')
    s, values = raw[sheets[0]], cached[sheets[0]]
    rows = [tuple(r) + (None,) * max(0, 23-len(r)) for r in s.values]
    cache = [tuple(r) + (None,) * max(0, 23-len(r)) for r in values.values]
    def resolved(row_index, col_index):
        value = rows[row_index][col_index]
        if not isinstance(value, str) or not value.startswith('='):
            return value
        value_cached = cache[row_index][col_index]
        if value_cached is not None:
            return value_cached
        # Resolve only a direct local-sheet reference, never evaluate formulas.
        match = re.fullmatch(r"=(?:'([^']+)'|([A-Za-z_][A-Za-z_0-9 ]*))!([A-Z]+[0-9]+)", value)
        if match:
            sheet = match.group(1) or match.group(2)
            if sheet in cached.sheetnames:
                return cached[sheet][match.group(3)].value
        raise ValueError(f'Fórmula de identificação sem valor disponível: LD!{openpyxl.utils.get_column_letter(col_index+1)}{row_index+1}')
    if len(rows) < 4:
        raw.close(); cached.close()
        raise ValueError('A aba LD está vazia ou não contém o cabeçalho esperado na linha 4.')
    roles = [str(clean_text(resolved(3,j))) for j in range(7,16)]
    if len(set(roles)) != 9 or any(x in ('None', '') for x in roles):
        raw.close(); cached.close()
        raise ValueError('Cabeçalho de funções incompatível com o layout LD')
    adjusted_layout = cache[3][21] == 'TOTAL Hh (ajustado)'
    records, wbs, issues, seen = [], [], [], set()
    current = None
    for i, (r, v) in enumerate(zip(rows, cache), 1):
        code = clean_text(resolved(i-1,1))
        eap, name = (resolved(i-1,j) for j in range(2,4))
        if isinstance(code, str):
            code = code.upper()
        if isinstance(code, str) and re.fullmatch(r'\d+(?:\.\d+)+', code):
            current = node(code, name, i)
            wbs.append(current)
            continue
        if code not in CODES or not clean_text(name):
            continue
        source = f'LD!{i}'
        semantic = json.dumps([current['codigo'] if current else None, code, eap, name], ensure_ascii=False)
        rid = (identities or {}).get(semantic) or stable_id('REG', semantic)
        if rid in seen:
            raise ValueError(f'Identidade duplicada em {source}; preencher identidades no JSON: {semantic}')
        seen.add(rid)
        rec = {'id': rid, 'wbs_id': current['id'] if current else None,
               'wbs_codigo': current['codigo'] if current else None,
               'codigo': code, 'eap_original': eap, 'nome': name,
               'tipo': 'documento' if code.startswith('D') else 'auxiliar',
               'origem': source, 'linha': i, 'unidade': r[5], 'quantidade': v[4],
               'a1_equivalente': None, 'hh': {}, 'total_original': None,
               'total_cache': v[17], 'total_ajustado_cache': v[21] if adjusted_layout else None,
               'fator_ajuste_original': r[19] if adjusted_layout else None, 'identidade_semantica': semantic,
               'originais': json.dumps(r, ensure_ascii=False, default=str),
               'cache': json.dumps(v, ensure_ascii=False, default=str)}
        try:
            # Only reproduce the documented A-format conversion when G has the
            # exact repository template formula. Unknown formulas use cached values.
            unit = clean_text(resolved(i-1,5))
            if isinstance(unit, str):
                unit = unit.upper()
            rec['unidade'] = unit
            expected = (f'=IF(F{i}="A4",E{i}/8,IF(F{i}="A3",E{i}/4,IF(F{i}="A2",E{i}/2,'
                        f'IF(F{i}="A1",E{i},IF(F{i}="A0",E{i}*2,IF(F{i}="HH","",""))))))')
            if unit == 'HH':
                multiplier = Decimal(1)
            elif unit in FORMATS and (r[6] == expected or not isinstance(r[6], str)):
                quantity = number(v[4])
                if quantity is None:
                    multiplier = None
                else:
                    multiplier = quantity * FORMATS[unit]
                    if number(v[6]) is not None and abs(number(v[6])-multiplier) > Decimal('0.000001'):
                        raise ValueError('Cache A1 divergente da fórmula/quantidade')
                rec['a1_equivalente'] = multiplier
            else:
                multiplier = number(v[6])
                rec['a1_equivalente'] = multiplier
                issues.append(issue('FORMULA_A1_NAO_PADRAO', source, 'A1 obtido do cache; revisar unidade/fórmula.'))
            inputs = [number(x) for x in v[7:16]]
            if any(isinstance(x, str) and x.startswith('=') for x in r[7:16]):
                issues.append(issue('FUNCAO_COM_FORMULA', source, 'HH por função dependem do cache Excel.'))
                if any(isinstance(r[j], str) and r[j].startswith('=') and v[j] is None for j in range(7,16)):
                    raise ValueError('Fórmula de HH sem cache; recalcular no Excel')
            if any(x is not None and x > 0 for x in inputs):
                if multiplier is None:
                    raise ValueError('Quantidade/A1 ausente para HH por A1')
                rec['hh'] = {role: value * multiplier for role, value in zip(roles, inputs) if value is not None and value > 0}
                rec['total_original'] = sum(rec['hh'].values(), Decimal(0))
                if number(v[17]) is not None and abs(number(v[17])-rec['total_original']) > Decimal('0.000001'):
                    issues.append(issue('TOTAL_DIVERGENTE', source, 'Soma por função diverge de R; registro bloqueado.'))
                    rec['hh'] = {}; rec['total_original'] = None
            else:
                if number(v[17]) is not None and number(v[17]) > 0:
                    raise ValueError('Total de HH preenchido, mas HH por função ausentes; revisar o mix na LD')
                issues.append(issue('SEM_HH', source, f'{name}: HH ausentes ou zero; não planejar horas.'))
        except ValueError as exc:
            rec['hh'] = {}; rec['total_original'] = None
            issues.append(issue('DADO_NAO_RESOLVIDO', source, str(exc)))
        if rec['tipo'] == 'documento' and ('reemiss' in str(name).lower() or 'atualizad' in str(name).lower()):
            issues.append(issue('CICLO_A_CONFIRMAR', source, 'Reemissão/atualização: ciclo de quatro etapas aplicado; validar aplicabilidade.'))
        records.append(rec)
    raw.close(); cached.close()
    if not records:
        raise ValueError('Nenhum documento ou atividade auxiliar reconhecido na aba LD. Confira o layout e os códigos DC/DB/DD/AC/AB/AD.')
    return records, wbs, issues
