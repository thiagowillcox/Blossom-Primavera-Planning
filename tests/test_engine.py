import copy
import json
from decimal import Decimal as D
from pathlib import Path
import openpyxl
import pytest
from blossom.ld import read_ld
from blossom.hours import distribute
from blossom.duration import duration
from blossom.relationships import build_links
from blossom.validation import topology, validate_config
from blossom.schedule import estimate
from blossom.__main__ import run
from blossom.common import number
from blossom.fixed_manual import get_fixed_manual

ROOT=Path(__file__).resolve().parents[1]

@pytest.fixture
def config():
    return json.loads((ROOT/'config/parametros.json').read_text(encoding='utf-8'))


def record(kind='documento'):
    return {'id':'doc-a','nome':'Desenho','tipo':kind,'wbs_id':'w',
            'origem':'LD!10','hh':{'EPleno':D(80),'PSênior':D(40)},'total_original':D(120)}


def test_distribution_by_role_and_stage(config):
    activities,assignments,audit,_=distribute([record()],config)
    assert [a['hh'] for a in activities]==list(map(D,['96','12','24','12']))
    for role,expected in [('EPleno',['64','8','16','8']),('PSênior',['32','4','8','4'])]:
        assert [a['hh'] for a in assignments if a['funcao']==role]==list(map(D,expected))
    assert all(r['diferenca']==0 for r in audit)
    assert sum(a['hh'] for a in activities)==D(144)


def test_auxiliary_no_increase(config):
    acts,attrs,audit,_=distribute([record('auxiliar')],config)
    assert len(acts)==1 and acts[0]['hh']==120
    assert all(r['hh_planejada']==r['hh_original'] for r in audit)
    assert build_links(acts,config)==[]


def test_parallel_duration_and_client_deadlines(config):
    config['capacidade_profissionais']={'EPleno':2,'PSênior':1}
    assert duration({'EPleno':D(80),'PSênior':D(40)},config,1)==5
    config['duracao_cliente_dias']['2']=7
    assert duration({'EPleno':D(8000)},config,2)==7
    config['capacidade_profissionais']={}
    assert duration({'EPleno':D(80)},config,1) is None


def test_links_and_schedule(config):
    config['capacidade_profissionais']={'EPleno':1,'PSênior':1}
    config['duracao_cliente_dias']={'2':5,'4':2}
    acts,_,_,_=distribute([record()],config)
    links=build_links(acts,config)
    assert len(links)==3
    estimate(acts,links)
    assert [(a['inicio_dia'],a['fim_dia']) for a in acts]==[(0,8),(8,13),(13,15),(15,17)]

@pytest.mark.parametrize('links',[
    [{'predecessora':'a','sucessora':'missing'}],
    [{'predecessora':'a','sucessora':'b'},{'predecessora':'b','sucessora':'a'}],
    [{'predecessora':'a','sucessora':'a'}],
    [{'predecessora':'a','sucessora':'b'},{'predecessora':'a','sucessora':'b'}],
])
def test_invalid_graphs(links):
    with pytest.raises(ValueError):topology([{'id':'a'},{'id':'b'}],links)


def test_duplicate_ids():
    with pytest.raises(ValueError):topology([{'id':'a'},{'id':'a'}],[])


def write_ld(path, reverse=False):
    w=openpyxl.Workbook();s=w.active;s.title='LD'
    for j,role in enumerate(['Gerente','Especialista','ESênior','EPleno','EJúnior','PSênior','PPleno','PJúnior','Desenh.'],8):s.cell(4,j,role)
    s.cell(7,2,'3.0');s.cell(7,4,'DETALHADA');s.cell(8,2,'3.1');s.cell(8,4,'MECÂNICA')
    rows=[['DD',None,'Documento A',2,'A1',2,None,None,None,8,None,None,None,None,None],['AD',None,'Apoio',None,'HH',None,None,None,None,16,None,None,None,None,None],['DD',None,'Sem horas',1,'A1',1]]
    if reverse:rows.reverse()
    for i,row in enumerate(rows,9):
        for j,v in enumerate(row,2):s.cell(i,j,v)
    w.save(path)


def test_ld_units_empty_and_stable_ids(tmp_path):
    path=tmp_path/'LD.xlsx';write_ld(path)
    records,wbs,issues=read_ld(path)
    assert records[0]['hh']=={'EPleno':D(16)}
    assert records[1]['hh']=={'EPleno':D(16)}
    assert not records[2]['hh']
    assert any(i['codigo']=='SEM_HH' for i in issues)
    ids={r['nome']:r['id'] for r in records}
    write_ld(path,True)
    reordered,_,_=read_ld(path)
    assert ids=={r['nome']:r['id'] for r in reordered}


@pytest.mark.parametrize('placeholder', [' ', '\u00a0', r'\xa0'])
def test_ld_blank_role_placeholders_preserve_hours(tmp_path, config, placeholder):
    path=tmp_path/'LD.xlsx';write_ld(path)
    workbook=openpyxl.load_workbook(path);sheet=workbook['LD']
    sheet['L9']=placeholder
    sheet['R9']=16
    sheet['H11']=placeholder
    workbook.save(path);workbook.close()
    records,_,issues=read_ld(path)
    assert records[0]['hh']=={'EPleno':D(16)}
    assert not records[2]['hh']
    assert not any(i['codigo']=='DADO_NAO_RESOLVIDO' for i in issues)
    activities,_,_,_=distribute(records,config)
    assert len(activities)==5
    assert sum(a['hh'] for a in activities)==D('35.2')


def test_ld_unknown_role_text_still_blocks_record(tmp_path):
    path=tmp_path/'LD.xlsx';write_ld(path)
    workbook=openpyxl.load_workbook(path)
    workbook['LD']['L9']='a confirmar'
    workbook.save(path);workbook.close()
    records,_,issues=read_ld(path)
    assert not records[0]['hh']
    assert any(i['codigo']=='DADO_NAO_RESOLVIDO' and i['origem']=='LD!9' for i in issues)


@pytest.mark.parametrize('value,expected', [
    ('12', '12'), (' 12,50 ', '12.5'), ('12.50', '12.5'),
    ('1.234,50', '1234.5'), ('1,234.50', '1234.5'),
    ('0.0625', '0.0625'), ('\ufeff\u200b', None),
])
def test_numeric_text_unambiguous(value, expected):
    assert number(value) == (D(expected) if expected is not None else None)


@pytest.mark.parametrize('value', ['1.234', '1,234', 'a confirmar', '#VALUE!', '-2', True, float('nan')])
def test_numeric_text_ambiguous_or_invalid(value):
    with pytest.raises(ValueError):
        number(value)


def test_ld_normalized_labels_and_numeric_text_preserve_sources(tmp_path):
    path=tmp_path/'LD.xlsx';write_ld(path)
    workbook=openpyxl.load_workbook(path);sheet=workbook['LD'];sheet.title=' ld '
    sheet['B9']=' dd ';sheet['D9']=' Documento A ';sheet['F9']=' a1 '
    sheet['E9']='2';sheet['K9']='8,50';sheet['R9']='17,00'
    sheet['H4']=' Gerente\u00a0'
    workbook.save(path);workbook.close()
    records,_,issues=read_ld(path)
    assert records[0]['hh']=={'EPleno':D(17)}
    assert records[0]['nome']==' Documento A '
    assert ' dd ' in records[0]['originais']
    assert not any(i['codigo']=='DADO_NAO_RESOLVIDO' for i in issues)


def test_ld_total_without_role_mix_is_flagged(tmp_path):
    path=tmp_path/'LD.xlsx';write_ld(path)
    workbook=openpyxl.load_workbook(path);workbook['LD']['R11']=20
    workbook.save(path);workbook.close()
    records,_,issues=read_ld(path)
    assert not records[2]['hh']
    assert any(i['codigo']=='DADO_NAO_RESOLVIDO' and i['origem']=='LD!11' for i in issues)


@pytest.mark.parametrize('mode', ['missing_sheet', 'empty_sheet', 'no_records', 'duplicate_roles'])
def test_ld_incompatible_layout_explained(tmp_path, mode):
    path=tmp_path/'LD.xlsx';write_ld(path)
    workbook=openpyxl.load_workbook(path);sheet=workbook['LD']
    if mode=='missing_sheet':sheet.title='Orçamento'
    elif mode=='empty_sheet':sheet.delete_rows(1,sheet.max_row)
    elif mode=='no_records':sheet.delete_rows(9,3)
    else:sheet['I4']=sheet['H4'].value
    workbook.save(path);workbook.close()
    with pytest.raises(ValueError,match='aba LD|cabeçalho|Nenhum|Cabeçalho'):
        read_ld(path)


def test_no_automatic_ambiguous_relationship(config):
    config['relacionamentos_aprovados']=[{'classe':'referencia_tecnica','aprovada':True,'predecessora':'x','sucessora':'y','regra_id':'r'}]
    assert build_links([],config)==[]

@pytest.mark.parametrize('change',[
    {'jornada_hh_dia':0}, {'capacidade_profissionais':{'EPleno':-1}},
    {'percentuais':[.8,.1,.1,.1]}, {'duracao_cliente_dias':{'2':-1}},
])
def test_bad_config(config,change):
    config.update(change)
    with pytest.raises(ValueError):validate_config(config)


def test_real_workbook_reproducibility(tmp_path):
    if not (ROOT/'input/LD.xlsm').exists():
        pytest.skip('Entradas comerciais locais não disponíveis')
    outputs=[tmp_path/'one',tmp_path/'two']
    for directory in outputs:
        tables=run(ROOT/'input/LD.xlsm',get_fixed_manual(),ROOT/'config/parametros.json',directory,ROOT/'config/observacoes_manual.json')
        assert len(tables['WBS'])==56
        assert all(r['codigo'][0].isdigit() for r in tables['WBS'])
        assert all('originais' not in r for r in tables['WBS'])
        original=json.loads((directory/'resultado.json').read_text(encoding='utf-8'))['registros']
        traced={r['registro_id'] for r in tables['Auditoria_HH']}
        traced.update(r['registro_id'] for r in tables['Pendencias'] if 'registro_id' in r)
        assert traced == {r['id'] for r in original}
        assert all('originais' in r and 'cache' in r for r in tables['Auditoria_HH'])
        assert all(r.get('predecessora') is None and r.get('sucessora') is None
                   for r in tables['Matriz_Regras'] if r['natureza']=='evidencia_nao_interpretada')
        assert len(tables['Atividades'])==952
        assert len(tables['Relacionamentos'])==655
        assert len([r for r in tables['Relacionamentos'] if r['origem'].startswith('REGRA-')])==16
        assert any(r['contexto_imagem']=='fluxo_disciplina' and r['disciplina']=='PROCESSO MINERAL'
                   for r in tables['Matriz_Regras'] if '!R227C2' in r.get('origem','') and 'FEL 4' in r.get('fase',''))
        assert abs(sum(a['hh_original'] for a in tables['Auditoria_HH'])-D('13442.25'))<D('0.000001')
        assert abs(sum(a['hh_planejada'] for a in tables['Auditoria_HH'])-D('15199.65'))<D('0.000001')
        assert any(p['origem']=='LD!790' and p['codigo']=='DADO_NAO_RESOLVIDO' for p in tables['Pendencias'])
        assert all(a['diferenca']==0 for a in tables['Auditoria_HH'])
        assert {r['classe'] for r in tables['Matriz_Regras']} >= {'condicional','referencia_tecnica','informacao_externa','a_validar'}
        w=openpyxl.load_workbook(directory/'pre_cronograma.xlsx',read_only=True)
        assert w['WBS'].max_row==57
        assert w.sheetnames==list(tables)
        assert len(w.sheetnames)==8
        w.close()
    for name in ('resultado.json','pre_cronograma.xlsx','matriz_regras.json','pendencias.json'):
        assert (outputs[0]/name).read_bytes()==(outputs[1]/name).read_bytes()


def test_mandatory_relationship_requires_known_rule(config):
    config['relacionamentos_aprovados']=[{
        'classe':'obrigatoria', 'aprovada':True, 'predecessora':'a',
        'sucessora':'b', 'regra_id':'inventada'}]
    with pytest.raises(ValueError, match='rastreável'):
        build_links([],config,[])


def test_validation_rejects_swapped_stage_hours(config):
    from blossom.validation import validate_result
    acts,assignments,audit,_=distribute([record()],config)
    tables={'WBS':[{'id':'w','pai':None}], 'Atividades':acts,
            'Atribuicoes_Funcoes':assignments, 'Auditoria_HH':audit,
            'Relacionamentos':build_links(acts,config), 'Matriz_Regras':[]}
    validate_result(tables)
    # Preserve the overall total, but corrupt the stage mix and its activity totals.
    assignments[0]['hh']-=1
    assignments[2]['hh']+=1
    acts[0]['hh']-=1
    acts[1]['hh']+=1
    with pytest.raises(ValueError, match='etapa/função'):
        validate_result(tables)
