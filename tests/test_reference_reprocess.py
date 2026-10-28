from io import BytesIO
from pathlib import Path
from decimal import Decimal
import pytest
import openpyxl
from blossom.reference import reprocess_reference

ROOT=Path(__file__).resolve().parents[1]

@pytest.fixture
def reference():
    p=ROOT/'input/pre_cronograma_v0001.xlsx'
    if not p.exists():pytest.skip('Referência comercial local não disponível')
    return p.read_bytes()


def workbook_bytes(w):
    b=BytesIO();w.save(b);w.close();return b.getvalue()


def test_every_reference_cell_is_preserved_or_recomputed_equal(reference):
    output,summary=reprocess_reference(reference)
    assert summary['alteracoes_relevantes']==0
    assert summary['atividades']==952
    assert summary['relacionamentos_ativos']==858
    assert abs(Decimal(summary['hh_planejadas'])-Decimal('14907.5'))<Decimal('0.00000001')
    a=openpyxl.load_workbook(BytesIO(reference),read_only=True,data_only=False)
    b=openpyxl.load_workbook(BytesIO(output),read_only=True,data_only=False)
    assert a.sheetnames==b.sheetnames
    for name in a.sheetnames:
        assert (a[name].max_row,a[name].max_column)==(b[name].max_row,b[name].max_column)
        for ra,rb in zip(a[name].values,b[name].values):
            for va,vb in zip(ra,rb):
                if isinstance(va,(int,float)) and not isinstance(va,bool) and isinstance(vb,(int,float)):
                    assert abs(Decimal(str(va))-Decimal(str(vb)))<Decimal('0.00000001'),(name,va,vb)
                else:assert va==vb,(name,va,vb)
    a.close();b.close()


def test_does_not_copy_precomputed_results(reference):
    w=openpyxl.load_workbook(BytesIO(reference));s=w['Atividades'];headers={c.value:c.column for c in s[1]}
    for field in ('hh','duracao_dias','inicio_dia_util','fim_dia_util'):s.cell(2,headers[field],9999)
    output,summary=reprocess_reference(workbook_bytes(w))
    assert summary['alteracoes_relevantes']>=4
    result=openpyxl.load_workbook(BytesIO(output),read_only=True)
    row=list(result['Atividades'].values)[1];values=dict(zip(headers,row))
    assert values['hh']==120
    assert values['duracao_dias']==15
    assert values['inicio_dia_util']==0
    assert values['fim_dia_util']==15
    result.close()


def test_rejects_missing_audit_cache(reference):
    w=openpyxl.load_workbook(BytesIO(reference));s=w['Auditoria_HH'];headers={c.value:c.column for c in s[1]}
    s.cell(2,headers['cache'],'[null]')
    with pytest.raises(ValueError,match='Cache de auditoria incompleto'):
        reprocess_reference(workbook_bytes(w))


def test_rejects_orphan_relationship(reference):
    w=openpyxl.load_workbook(BytesIO(reference));s=w['Relacionamentos'];headers={c.value:c.column for c in s[1]}
    s.cell(2,headers['sucessora'],'inexistente')
    with pytest.raises(ValueError,match='atividade inexistente'):
        reprocess_reference(workbook_bytes(w))


def test_reference_can_be_processed_through_site(reference):
    from streamlit.testing.v1 import AppTest
    app=AppTest.from_file(str(ROOT/'app.py')).run(timeout=30)
    app.session_state['reference_file']=('pre_cronograma_v0001.xlsx',reference)
    app.radio[0].set_value('Configuração').run()
    app.selectbox[0].set_value('Reprocessar referência Wardi').run()
    assert not app.exception
    assert not app.button[0].disabled
    app.button[0].click().run(timeout=60)
    assert not app.exception
    assert app.radio[0].value=='Resultados'
    assert any('0 diferenças relevantes' in s.value for s in app.success)
