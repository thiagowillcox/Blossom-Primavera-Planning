import hashlib
import json
from collections import Counter
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from streamlit.testing.v1 import AppTest

from blossom.decisions import WBSDecisionRequired
from blossom.ld import read_ld
from blossom.hours import distribute
from blossom.web import inspect_wbs

ROOT=Path(__file__).resolve().parents[1]


@pytest.fixture
def duplicate_ld(tmp_path):
    path=tmp_path/'LD.xlsx'
    workbook=openpyxl.Workbook();sheet=workbook.active;sheet.title='LD'
    for j,role in enumerate(['Gerente','Especialista','ESênior','EPleno','EJúnior','PSênior','PPleno','PJúnior','Desenh.'],8):
        sheet.cell(4,j,role)
    for row,code,name in [(7,'4.0','Complementares'),(8,'4.1','SCAN 3D'),(10,'4.1','Hazid/Hazop'),(12,'4.2','As built')]:
        sheet.cell(row,2,code);sheet.cell(row,4,name)
    for row,name in [(9,'Documento SCAN'),(11,'Documento Hazop'),(13,'Documento As built')]:
        for col,value in {2:'DD',4:name,5:1,6:'A1',7:1,11:10,18:10}.items():
            sheet.cell(row,col,value)
    workbook.save(path);workbook.close()
    return path


def choices(path, code='4.5'):
    return {'ld_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'codigos_por_origem':{'LD!B10:D10':code}}


def test_duplicate_requires_user_decision(duplicate_ld):
    with pytest.raises(WBSDecisionRequired) as error:
        read_ld(duplicate_ld)
    assert [n['nome'] for n in error.value.groups['4.1']]==['SCAN 3D','Hazid/Hazop']
    assert inspect_wbs(duplicate_ld.read_bytes())==error.value.groups


def test_decision_preserves_source_effort_and_hierarchy(duplicate_ld):
    original=duplicate_ld.read_bytes()
    config=json.loads((ROOT/'config/parametros.json').read_text())
    adjustments=choices(duplicate_ld)
    records,nodes,issues=read_ld(duplicate_ld,wbs_adjustments=adjustments)
    assert duplicate_ld.read_bytes()==original
    assert [r['wbs_codigo'] for r in records]==['4.1','4.5','4.2']
    assert len({n['id'] for n in nodes})==4
    assert nodes[2]['codigo_original']=='4.1'
    assert 'DECISAO_WBS_APLICADA' in [i['codigo'] for i in issues]
    activities,_,_,_=distribute(records,config)
    assert len(activities)==12
    assert sum(a['hh'] for a in activities)==Decimal(36)
    assert Counter(a['wbs_id'] for a in activities)==Counter({n['id']:4 for n in nodes[1:]})
    assert inspect_wbs(original,adjustments)=={}


@pytest.mark.parametrize('code,message', [('4.2','repetidos'),('9.1','pai'),('abc','inválido')])
def test_invalid_decisions_stay_blocked(duplicate_ld,code,message):
    with pytest.raises(ValueError,match=message):
        read_ld(duplicate_ld,wbs_adjustments=choices(duplicate_ld,code))


def test_decisions_bound_to_source_and_origin(duplicate_ld):
    adjustment=choices(duplicate_ld);adjustment['ld_sha256']='other-file'
    with pytest.raises(ValueError,match='outra LD'):
        read_ld(duplicate_ld,wbs_adjustments=adjustment)
    adjustment=choices(duplicate_ld);adjustment['codigos_por_origem']={'LD!B999:D999':'4.5'}
    with pytest.raises(ValueError,match='inexistente'):
        read_ld(duplicate_ld,wbs_adjustments=adjustment)


def test_renaming_branch_does_not_reassign_children_silently(duplicate_ld):
    workbook=openpyxl.load_workbook(duplicate_ld)
    workbook['LD']['B11']='4.1.1';workbook['LD']['D11']='Subgrupo Hazop'
    workbook.save(duplicate_ld);workbook.close()
    adjustment=choices(duplicate_ld)
    with pytest.raises(ValueError,match='subgrupo'):
        read_ld(duplicate_ld,wbs_adjustments=adjustment)
    adjustment['codigos_por_origem']['LD!B11:D11']='4.5.1'
    _,nodes,_=read_ld(duplicate_ld,wbs_adjustments=adjustment)
    assert nodes[3]['pai']==nodes[2]['id']


def test_site_requests_choice_and_invalidates_on_new_ld(duplicate_ld):
    app=AppTest.from_file(str(ROOT/'app.py'))
    app.session_state['ld_file']=('LD.xlsx',duplicate_ld.read_bytes())
    app.session_state['manual_file']=('Manual.xlsx',b'not-needed-for-preflight')
    app.session_state['navigation']='Configuração'
    app.run(timeout=30)
    assert not app.exception
    assert any('Precisamos da sua escolha' in s.value for s in app.subheader)
    assert next(b for b in app.button if b.label.startswith('Gerar')).disabled
    fields=[t for t in app.text_input if t.label.startswith('Hazid')]
    fields[0].set_value('4.5')
    next(b for b in app.button if b.label=='Confirmar códigos e continuar').click().run(timeout=30)
    assert not app.exception
    assert not next(b for b in app.button if b.label.startswith('Gerar')).disabled
    assert app.session_state['wbs_decisions']['codigos_por_origem']['LD!B10:D10']=='4.5'
    app.radio[0].set_value('Arquivos').run()
    app.radio[0].set_value('Configuração').run()
    assert not next(b for b in app.button if b.label.startswith('Gerar')).disabled
    # Changing the file, even to another revision of the same LD, requires review.
    workbook=openpyxl.load_workbook(duplicate_ld);workbook['LD']['D7']='Outro projeto'
    workbook.save(duplicate_ld);workbook.close()
    app.session_state['ld_file']=('Outro.xlsx',duplicate_ld.read_bytes())
    app.run(timeout=30)
    assert next(b for b in app.button if b.label.startswith('Gerar')).disabled

