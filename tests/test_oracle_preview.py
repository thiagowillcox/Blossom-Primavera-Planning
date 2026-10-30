import json
from pathlib import Path
import openpyxl
import pytest
from blossom.oracle.preview import prepare_preview
from blossom.oracle.client import ConnectionFailure


def make_files(tmp_path, duration=None):
    data={'WBS':[{'id':'w','pai':None}],
          'Atividades':[{'id':'a','registro_id':'r','tipo_linha':'atividade','etapa':0,'wbs_id':'w','hh':10,'duracao_dias':duration}],
          'Atribuicoes_Funcoes':[{'atividade_id':'a','hh':10}],
          'Relacionamentos':[], 'Matriz_Regras':[],
          'Parametros':[{'parametro':'ld_sha256','valor':'source_hash'}],
          'Pendencias':[{'codigo':'DURACAO_PENDENTE'}],
          'Auditoria_HH':[{'diferenca':0}]}
    workbook=openpyxl.Workbook();workbook.remove(workbook.active)
    for name,rows in data.items():
        sheet=workbook.create_sheet(name);headers=list(rows[0]) if rows else ['status'];sheet.append(headers)
        for row in rows:sheet.append([row.get(k) for k in headers])
    path=tmp_path/'schedule.xlsx';workbook.save(path);workbook.close()
    selection=tmp_path/'selection.json';selection.write_text(json.dumps({'servidor':'https://primavera-us2.oraclecloud.com',
        'projeto_selecionado':{'projectId':7,'projectName':'Laboratório'}}))
    return path,selection,tmp_path/'preview.json'


def test_offline_preview_never_claims_upload_ready(tmp_path):
    path,selection,destination=make_files(tmp_path)
    report=prepare_preview(path,selection,destination)
    assert report['contagens']['atividades']==1
    assert report['hh_planejadas']=='10'
    assert report['atividades_sem_duracao']==1
    assert report['pronto_para_envio'] is False
    assert report['envio_implementado'] is False
    assert json.loads(destination.read_text())==report


def test_known_durations_do_not_bypass_oracle_schema_mapping(tmp_path):
    report=prepare_preview(*make_files(tmp_path,duration=1))
    assert report['atividades_sem_duracao']==0 and report['pronto_para_envio'] is False


def test_invalid_or_foreign_project_selection_is_blocked(tmp_path):
    path,selection,destination=make_files(tmp_path)
    selection.write_text(json.dumps({'servidor':'https://other.example.org','projeto_selecionado':{'projectId':7}}))
    with pytest.raises(ConnectionFailure):prepare_preview(path,selection,destination)
    assert not destination.exists()


def test_changed_assignment_hours_are_detected(tmp_path):
    path,selection,destination=make_files(tmp_path)
    workbook=openpyxl.load_workbook(path);workbook['Atribuicoes_Funcoes']['B2']=11
    workbook.save(path);workbook.close()
    with pytest.raises(ConnectionFailure,match='HH'):prepare_preview(path,selection,destination)
    assert not destination.exists()
