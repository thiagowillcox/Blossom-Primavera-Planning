import json
from pathlib import Path
from decimal import Decimal as D
import pytest
import openpyxl
from blossom.ld import read_ld
from blossom.duration import duration
from blossom.validation import validate_config
from blossom.wbs import node

ROLES=['Gerente','Especialista','ESênior','EPleno','EJúnior','PSênior','PPleno','PJúnior','Desenh.']


def workbook(path, unresolved=False):
    w=openpyxl.Workbook();s=w.active;s.title='LD';cpu=w.create_sheet('CPU ')
    for i,role in enumerate(ROLES):
        col=openpyxl.utils.get_column_letter(5+i)
        cpu.cell(2,5+i,role)
        s.cell(4,8+i,f"='CPU '!{col}2")
    if unresolved:s['H4']="='Ausente'!E2"
    s['B7']='3.0';s['D7']='DETALHADA';s['B8']='3.1';s['D8']='MECÂNICA'
    s['B9']='DD';s['D9']='Documento A';s['E9']=2;s['F9']='A1';s['G9']=2;s['K9']=8;s['R9']=16
    s['T9']=2000;s['V9']=3000;s['V4']='Especialista'
    w.save(path)


def test_local_formula_headers_and_costs(tmp_path):
    p=tmp_path/'budget.xlsx';workbook(p)
    records,wbs,issues=read_ld(p)
    assert records[0]['hh']=={'EPleno':D(16)}
    assert records[0]['total_original']==16
    assert records[0]['total_ajustado_cache'] is None
    assert records[0]['fator_ajuste_original'] is None
    assert '3000' in records[0]['cache']


def test_unresolved_role_header_is_rejected(tmp_path):
    p=tmp_path/'budget.xlsx';workbook(p,True)
    with pytest.raises(ValueError,match='identificação sem valor'):
        read_ld(p)


def test_nested_wbs_parent():
    assert node('3.1.2','Grupo',10)['pai']==node('3.1','Disciplina',9)['id']
    assert node('3.1.0','Grupo',10)['pai']==node('3.1','Disciplina',9)['id']



def test_effort_is_not_elapsed_time():
    c={'jornada_hh_dia':8,'capacidade_profissionais':{'EPleno':1,'PSênior':1},'duracao_cliente_dias':{'2':3,'4':3}}
    assert duration({'EPleno':D(8),'PSênior':D(8)},c,1)==1


def test_old_effort_only_mode_is_rejected():
    c=json.loads(Path('config/parametros.json').read_text(encoding='utf-8'));c['metodo_duracao']='equivalente_hh'
    with pytest.raises(ValueError,match='capacidade_funcoes'):validate_config(c)


def test_reconciliation_uses_source_total_not_reference():
    from blossom.source_audit import reconcile_ld
    records=[{'id':'r','nome':'Documento','origem':'LD!10','total_cache':D(100)}]
    audit=[{'registro_id':'r','hh_original':D(60)},{'registro_id':'r','hh_original':D(40)}]
    assert reconcile_ld(records,audit)[0]['conferencia_total_ld']=='conferido'
    audit[1]['hh_original']=D(39)
    result=reconcile_ld(records,audit)[0]
    assert result['diferenca_total_ld']==-1
    assert result['conferencia_total_ld']=='divergente'


def test_metadata_only_change_does_not_change_hours(tmp_path):
    from zipfile import ZipFile
    from blossom.common import digest
    a=tmp_path/'original.xlsx';b=tmp_path/'outro.xlsx';workbook(a)
    with ZipFile(a) as source,ZipFile(b,'w') as target:
        for info in source.infolist():
            data=source.read(info.filename)
            if info.filename=='docProps/core.xml':data=data.replace(b'</cp:coreProperties>',b'<!-- Apenas metadados -->\n</cp:coreProperties>')
            target.writestr(info,data)
    assert digest(a)!=digest(b)
    before,_,_=read_ld(a);after,_,_=read_ld(b)
    assert before[0]['hh']==after[0]['hh']
    assert before[0]['total_original']==after[0]['total_original']
