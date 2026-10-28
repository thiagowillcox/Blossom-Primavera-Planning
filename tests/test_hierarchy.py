from decimal import Decimal
from pathlib import Path
import json
import openpyxl
import pytest
from blossom.hierarchy import activity_hierarchy
from blossom.export import save
from blossom.hours import distribute


def sample():
    config=json.loads(Path('config/parametros.json').read_text(encoding='utf-8'))
    records=[{'id':'doc','wbs_id':'discipline','nome':'Documento A','tipo':'documento',
              'hh':{'EPleno':Decimal(100)},'total_original':Decimal(100),'origem':'LD!10'},
             {'id':'aux','wbs_id':'discipline','nome':'Apoio','tipo':'auxiliar',
              'hh':{'EPleno':Decimal(8)},'total_original':Decimal(8),'origem':'LD!11'}]
    activities,assignments,audit,_=distribute(records,config)
    return {'WBS':[{'id':'phase','pai':None,'codigo':'3.0','nome':'DETALHADA','origem':'LD!7'},
                   {'id':'discipline','pai':'phase','codigo':'3.1','nome':'MECÂNICA','origem':'LD!8'}],
            'Atividades':activities,'Atribuicoes_Funcoes':assignments,'Relacionamentos':[],
            'Matriz_Regras':[],'Parametros':[],'Pendencias':[],'Auditoria_HH':audit}


def test_outline_preserves_leaves_and_hours():
    tables=sample()
    rows=activity_hierarchy(tables)
    assert [r['tipo_linha'] for r in rows]==['fase','agrupador','documento']+['atividade']*5
    assert [r['nivel'] for r in rows]==[0,1,2,3,3,3,3,2]
    assert rows[2]['nome']=='Documento A'
    assert all(r['pai_id']=='doc' for r in rows[3:7])
    assert rows[-1]['nome']=='Apoio' and rows[-1]['pai_id']=='discipline'
    assert {r['id'] for r in rows if r['tipo_linha']=='atividade'}=={r['id'] for r in tables['Atividades']}
    assert sum(r['hh'] or 0 for r in rows)==Decimal(128)
    assert all(r['hh'] is None for r in rows if r['tipo_linha']!='atividade')
    assert len({r['id'] for r in rows})==len(rows)


def test_excel_has_real_outline_and_indentation(tmp_path):
    tables=sample();save(tables,[],{},tmp_path)
    w=openpyxl.load_workbook(tmp_path/'pre_cronograma.xlsx')
    sheet=w['Atividades'];keys=[c.value for c in sheet[1]]
    name=keys.index('nome')+1
    assert sheet.sheet_properties.outlinePr.summaryBelow is False
    assert sheet.auto_filter.ref is None
    assert sheet.cell(5,name).value=='Emissão Inicial'
    assert sheet.cell(5,name).alignment.indent==3
    assert sheet.row_dimensions[5].outlineLevel==3
    assert sheet.row_dimensions[4].outlineLevel==2
    assert w['WBS'].max_row==3
    payload=json.loads((tmp_path/'resultado.json').read_text(encoding='utf-8'))
    assert len(payload['linhas_cronograma'])==8
    assert len(payload['tabelas']['Atividades'])==5
    w.close()


def test_unknown_parent_is_rejected():
    tables=sample();tables['WBS'][1]['pai']='missing'
    with pytest.raises(ValueError,match='pai inexistente'):
        activity_hierarchy(tables)
